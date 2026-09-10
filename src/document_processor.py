"""
Document Processing Pipeline.
Handles loading, chunking, and preparing documents for embedding.
Supports CSV, Excel, PDF, and TXT files.
"""

import os
import tempfile
from typing import Optional

import pandas as pd
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_community.document_loaders import (
    CSVLoader,
    PyPDFLoader,
    TextLoader,
)
import openpyxl

import config
from src.utils import get_file_extension


class DocumentProcessor:
    """Processes uploaded files into LangChain Document chunks."""

    def __init__(
        self,
        chunk_size: int = config.CHUNK_SIZE,
        chunk_overlap: int = config.CHUNK_OVERLAP,
    ):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    def load_file(self, file_path: str, original_filename: str) -> list[Document]:
        """
        Load a file and return a list of LangChain Documents.

        Args:
            file_path: Path to the temporary file on disk.
            original_filename: Original name of the uploaded file.

        Returns:
            List of Document objects with content and metadata.
        """
        ext = get_file_extension(original_filename)
        loader = self._get_loader(file_path, ext)
        documents = loader.load()

        # Enrich metadata
        for doc in documents:
            doc.metadata["source"] = original_filename
            doc.metadata["file_type"] = ext

        return documents

    def _get_loader(self, file_path: str, ext: str):
        """Select the appropriate document loader based on file extension."""
        if ext == "csv":
            return CSVLoader(file_path=file_path, encoding="utf-8")
        elif ext in ("xlsx", "xls"):
            return _ExcelLoader(file_path=file_path)
        elif ext == "pdf":
            return PyPDFLoader(file_path=file_path)
        elif ext == "txt":
            return TextLoader(file_path=file_path, encoding="utf-8")
        else:
            raise ValueError(f"Unsupported file type: .{ext}")

    def chunk_documents(self, documents: list[Document]) -> list[Document]:
        """
        Split documents into smaller chunks for embedding.

        Args:
            documents: List of loaded Document objects.

        Returns:
            List of chunked Document objects.
        """
        chunks = self.text_splitter.split_documents(documents)

        # Add chunk index to metadata
        for i, chunk in enumerate(chunks):
            chunk.metadata["chunk_index"] = i
            chunk.metadata["total_chunks"] = len(chunks)

        return chunks

    def process_file(self, file_path: str, original_filename: str) -> list[Document]:
        """
        Full pipeline: load → chunk a file.

        Args:
            file_path: Path to the file on disk.
            original_filename: Original name of the uploaded file.

        Returns:
            List of chunked Document objects ready for embedding.
        """
        documents = self.load_file(file_path, original_filename)
        chunks = self.chunk_documents(documents)
        return chunks

    @staticmethod
    def save_uploaded_file(uploaded_file) -> str:
        """
        Save a Streamlit UploadedFile to a temporary location.

        Args:
            uploaded_file: Streamlit UploadedFile object.

        Returns:
            Path to the saved temporary file.
        """
        ext = get_file_extension(uploaded_file.name)
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
            tmp.write(uploaded_file.getbuffer())
            return tmp.name

    @staticmethod
    def load_dataframe(file_path: str, original_filename: str) -> Optional[pd.DataFrame]:
        """
        Load CSV or Excel files as a Pandas DataFrame for direct analysis.

        Args:
            file_path: Path to the file.
            original_filename: Original filename to detect type.

        Returns:
            DataFrame if file is tabular, None otherwise.
        """
        ext = get_file_extension(original_filename)
        try:
            if ext == "csv":
                return pd.read_csv(file_path)
            elif ext in ("xlsx", "xls"):
                return pd.read_excel(file_path)
        except Exception:
            return None
        return None

    @staticmethod
    def cleanup_temp_file(file_path: str):
        """Remove a temporary file from disk."""
        try:
            if os.path.exists(file_path):
                os.unlink(file_path)
        except OSError:
            pass


class _ExcelLoader:
    """
    Lightweight Excel loader using openpyxl.
    Avoids the heavy unstructured/transformers dependency chain.
    """

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> list[Document]:
        """Load Excel file and return a list of Documents (one per sheet)."""
        documents = []
        wb = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                continue

            # First row as headers
            headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
            lines = [", ".join(headers)]

            for row in rows[1:]:
                line = ", ".join(str(cell) if cell is not None else "" for cell in row)
                lines.append(line)

            content = "\n".join(lines)
            documents.append(
                Document(
                    page_content=content,
                    metadata={"source": self.file_path, "sheet": sheet_name},
                )
            )

        wb.close()
        return documents
